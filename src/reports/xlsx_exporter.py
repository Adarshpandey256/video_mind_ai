"""Excel export: one workbook, one worksheet per selected content type.

Only worksheets that were selected and actually have data are created, so a
short recording does not produce a dozen empty tabs. Long text is wrapped,
headers are frozen and every user-controlled cell is written as text (never as
a formula), which removes the spreadsheet-injection risk without altering the
report model.

``openpyxl`` is imported lazily so the rest of the application keeps working if
the optional document libraries are not installed.
"""

from __future__ import annotations

import io
from typing import Any

from src.reports.export_types import RenderResult
from src.reports.report_models import ReportSection, VideoMindReport
from src.utils.time_utils import format_timestamp

_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_WIDTHS = {"A": 28, "B": 60, "C": 32, "D": 40, "E": 32, "F": 28, "G": 22}


class WorkbookError(RuntimeError):
    """Raised when an Excel workbook cannot be produced."""


def _openpyxl():
    try:
        import openpyxl  # noqa: PLC0415 - optional dependency, imported on demand
    except ImportError as error:  # pragma: no cover - depends on the environment
        raise WorkbookError(
            "Excel export needs the openpyxl package. Install it with "
            "`pip install openpyxl` and restart the application."
        ) from error
    return openpyxl


def _set_text(cell: Any, value: object) -> None:
    """Write a literal string; a leading '=' must never become a formula."""
    if value is None:
        cell.value = ""
        return
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.value = value
        return
    text = str(value)
    cell.value = text
    if text.startswith(_DANGEROUS_PREFIXES):
        cell.data_type = "s"


def _write_table(
    worksheet: Any,
    headers: list[str],
    rows: list[list[object]],
    *,
    wrap_columns: tuple[int, ...] = (),
) -> None:
    letters = list(_openpyxl().utils.get_column_letter(index) for index in range(1, len(headers) + 1))
    for column, header in enumerate(headers, start=1):
        _set_text(worksheet.cell(row=1, column=column), header)
        letter = letters[column - 1]
        worksheet.column_dimensions[letter].width = _WIDTHS.get(letter, 28)
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_index, column=column)
            _set_text(cell, value)
            if column in wrap_columns:
                cell.alignment = _openpyxl().styles.Alignment(wrap_text=True, vertical="top")
    worksheet.freeze_panes = "A2"
    if headers:
        for cell in worksheet[1]:
            cell.font = _openpyxl().styles.Font(bold=True)


def _overview(report: VideoMindReport) -> tuple[list[str], list[list[object]]]:
    rows: list[list[object]] = [
        ["Report title", report.title],
        ["Report id", report.report_id],
        ["Generated at", report.generated_at.isoformat()],
        ["Schema version", report.schema_version],
        ["Author", report.author or ""],
        ["Organisation", report.organisation or ""],
        ["Video title", report.video.title],
        ["Video id", report.video.video_id],
        ["Source type", report.video.source_type.value],
        ["Source URL", report.video.source_url or ""],
        ["Original file", report.video.original_filename or ""],
        ["Duration", report.video.duration_label],
        ["Language", report.video.language or ""],
        ["Recorded", report.video.created_at.isoformat() if report.video.created_at else ""],
        ["Transcript segments", report.video.segment_count],
        ["Timestamps available", "yes" if report.timestamps_available else "no"],
        ["Summary style", report.summary.style if report.summary else ""],
        ["Chat answers included", report.chat_answers],
        ["Sections", ", ".join(section.value for section in report.sections)],
    ]
    if report.notes:
        rows.append(["Notes", report.notes])
    for entry in report.omitted_sections:
        rows.append([f"Not included: {entry.section.value}", entry.reason or "no data available"])
    return ["Field", "Value"], rows


def build_xlsx(report: VideoMindReport) -> RenderResult:
    """Render the report as an .xlsx workbook."""
    openpyxl = _openpyxl()
    workbook = openpyxl.Workbook()
    overview = workbook.active
    overview.title = "Overview"
    headers, rows = _overview(report)
    _write_table(overview, headers, rows, wrap_columns=(2,))

    if report.includes(ReportSection.TRANSCRIPT) and report.transcript:
        sheet = workbook.create_sheet("Transcript")
        sheet_headers = [
            "Segment",
            "Start (s)",
            "End (s)",
            "Start",
            "End",
            "Text",
        ]
        sheet_rows = [
            [
                entry.segment_id,
                ""
                if entry.start_seconds is None or not report.include_timestamps
                else round(entry.start_seconds, 3),
                ""
                if entry.end_seconds is None or not report.include_timestamps
                else round(entry.end_seconds, 3),
                ""
                if entry.start_seconds is None or not report.include_timestamps
                else format_timestamp(entry.start_seconds),
                ""
                if entry.end_seconds is None or not report.include_timestamps
                else format_timestamp(entry.end_seconds),
                entry.text,
            ]
            for entry in report.transcript
        ]
        _write_table(sheet, sheet_headers, sheet_rows, wrap_columns=(6,))
        sheet.column_dimensions["F"].width = 100

    if report.includes(ReportSection.KEY_TAKEAWAYS) and report.key_takeaways:
        sheet = workbook.create_sheet("Key takeaways")
        _write_table(
            sheet,
            ["Takeaway", "Explanation", "Timestamp", "Transcript excerpt"],
            [
                [item.title, item.explanation, item.time_label, item.excerpt or ""]
                for item in report.key_takeaways
            ],
            wrap_columns=(2, 4),
        )

    if report.includes(ReportSection.TOPICS) and report.topics:
        sheet = workbook.create_sheet("Topics")
        _write_table(
            sheet,
            ["Topic", "Explanation", "Key points", "Timestamp"],
            [
                [
                    item.title,
                    item.explanation,
                    "\n".join(f"- {point}" for point in item.key_points),
                    item.time_label,
                ]
                for item in report.topics
            ],
            wrap_columns=(2, 3),
        )

    if report.includes(ReportSection.CHAPTERS) and report.chapters:
        sheet = workbook.create_sheet("Chapters")
        _write_table(
            sheet,
            ["Chapter", "Description", "Major points", "Timestamp"],
            [
                [
                    item.title,
                    item.description,
                    "\n".join(f"- {point}" for point in item.major_points),
                    item.time_label,
                ]
                for item in report.chapters
            ],
            wrap_columns=(2, 3),
        )

    if report.includes(ReportSection.ACTION_ITEMS) and report.action_items:
        sheet = workbook.create_sheet("Actions")
        _write_table(
            sheet,
            ["Kind", "Text", "Owner", "Deadline", "Timestamp"],
            [
                [item.kind, item.text, item.owner or "", item.deadline or "", item.time_label]
                for item in report.action_items
            ],
            wrap_columns=(2,),
        )

    if report.includes(ReportSection.CHAT_QA) and report.chat:
        sheet = workbook.create_sheet("Chat Q&A")
        _write_table(
            sheet,
            ["Question", "Answer", "Cited sources", "Timestamps"],
            [
                [
                    entry.question,
                    entry.answer,
                    "; ".join(citation.source_ref for citation in entry.citations),
                    "; ".join(citation.time_label for citation in entry.citations),
                ]
                for entry in report.chat
            ],
            wrap_columns=(1, 2),
        )

    if report.includes(ReportSection.SOURCES) and report.sources:
        sheet = workbook.create_sheet("Sources")
        _write_table(
            sheet,
            ["Reference", "Source", "Type", "Timestamp", "Excerpt"],
            [
                [
                    source.source_ref,
                    source.label,
                    source.source_kind,
                    source.time_label,
                    source.excerpt,
                ]
                for source in report.sources
            ],
            wrap_columns=(2, 5),
        )

    buffer = io.BytesIO()
    try:
        workbook.save(buffer)
    except Exception as error:  # noqa: BLE001 - translated into an export error
        raise WorkbookError(
            "The Excel workbook could not be written. Try a smaller report or another format."
        ) from error
    return RenderResult(data=buffer.getvalue())


__all__ = ["WorkbookError", "build_xlsx"]
